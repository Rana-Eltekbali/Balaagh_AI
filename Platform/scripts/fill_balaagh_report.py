from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor


SOURCE = Path("attached_assets/SIC_AI_Capstone_Project_Final_Report_1790208057462.docx")
OUTPUT = Path("deliverables/SIC_AI_Capstone_Project_Balaagh_UI_Completed.docx")
ASSETS = Path("deliverables/report-assets")


def add_before(anchor, kind, value="", level=0):
    paragraph = anchor.insert_paragraph_before()
    if kind == "image":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        paragraph.add_run().add_picture(str(ASSETS / value), width=Inches(6.35))
        return paragraph

    if kind == "caption":
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        run.italic = True
        run.font.size = Pt(9)
        run.font.color.rgb = RGBColor(90, 102, 118)
        return paragraph

    if kind == "bullet":
        value = f"• {value}"

    run = paragraph.add_run(value)
    if kind == "heading":
        run.bold = True
        run.font.size = Pt(16 if level == 2 else 13)
        run.font.color.rgb = RGBColor(31, 61, 105)
        paragraph.paragraph_format.space_before = Pt(8)
        paragraph.paragraph_format.space_after = Pt(4)
    else:
        run.font.size = Pt(10.5)
        paragraph.paragraph_format.space_after = Pt(5)
    return paragraph


def main():
    doc = Document(SOURCE)
    anchor = next(
        paragraph
        for paragraph in doc.paragraphs
        if paragraph.text.strip() == "3.5. Testing and Improvements"
    )

    blocks = [
        (
            "text",
            "Balaagh AI was designed as a responsive command center for reviewing Arabic crisis reports in the Libyan context. The interface keeps the original Arabic report visible while converting it into a structured, reviewable record. The dashboard is intentionally a decision-support prototype: it does not dispatch emergency services, contact authorities, or claim real-time prediction.",
        ),
        ("heading", "Dashboard / Command Center", 3),
        (
            "text",
            "The dashboard gives the analyst an immediate overview of the saved demo dataset and the current review workload. It combines four headline indicators with recent reports and distribution views so that the most important patterns can be scanned before opening a detailed record.",
        ),
        ("bullet", "Total reports: the number of saved analyses in the current dataset."),
        ("bullet", "Critical reports and High priority: reports requiring focused human review."),
        ("bullet", "Reports today: a time-based activity indicator for the current day."),
        ("bullet", "Recent reports: a direct path to the latest structured records."),
        ("bullet", "Priority mix and incident classes: compact visual summaries of the saved reports."),
        ("image", "dashboard.jpg"),
        (
            "caption",
            "Figure 3.4.1. Balaagh AI Command Center dashboard with live demo data, recent reports, priority mix, and incident-class distribution.",
        ),
        ("heading", "Analyze Report Workspace", 3),
        (
            "text",
            "The Analyze Report screen is the main user flow. The analyst pastes an Arabic report into an RTL-ready input area, submits it for analysis, reviews the structured output, and saves the reviewed result. The screen shows a clear prototype disclaimer and keeps the original text available for verification.",
        ),
        (
            "bullet",
            "Arabic RTL input with an example report: “مثال: في حريق في منزل في جنزور وفي طفلين داخل البيت”.",
        ),
        (
            "bullet",
            "Structured output for incident class, priority, location, people at risk, required support, relevance, and summary.",
        ),
        (
            "bullet",
            "Loading, error, empty, and saved states so the workflow remains understandable during API requests.",
        ),
        ("image", "analyze.jpg"),
        ("caption", "Figure 3.4.2. RTL Arabic report input and structured-review workspace."),
        ("heading", "Reports, Locations, and Analytics", 3),
        (
            "text",
            "Supporting screens extend the dashboard into a small review system. Saved Reports provides search, sorting, filters, detail views, and deletion. Locations summarizes reporting areas and potential risk indicators. Analytics compares incident classes, support needs, priorities, locations, and illustrative model-evaluation values.",
        ),
        ("image", "analytics.jpg"),
        (
            "caption",
            "Figure 3.4.3. Analytics screen showing incident distribution, support distribution, and high-level metrics.",
        ),
        ("heading", "AI Approach Represented in the UI", 3),
        (
            "text",
            "The current implementation uses a transparent keyword-based Arabic baseline so that every result can be explained during the capstone demonstration. The modular design leaves a clear path for a trained NLP model: MARBERT or AraBERT can be fine-tuned for incident classification, Arabic named-entity recognition can extract locations and people, and a separate rule layer can combine model signals with priority and support checks. The interface labels evaluation values as demo or placeholder results and does not present them as production performance.",
        ),
        ("heading", "Usability, Accessibility, and Responsible Use", 3),
        ("bullet", "Responsive layout with a collapsible navigation menu for smaller screens."),
        ("bullet", "Clear hierarchy, color-coded priority badges, readable status indicators, and consistent loading and empty states."),
        ("bullet", "Arabic text is preserved in RTL presentation while the surrounding dashboard remains easy to scan in English."),
        (
            "bullet",
            "Human verification is required before treating any structured result as meaningful. The prototype is not emergency dispatch software and should not replace local expertise, source verification, or safety protocols.",
        ),
    ]

    for kind, value, *level in blocks:
        add_before(anchor, kind, value, level[0] if level else 0)

    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "3.4. User Interface":
            for run in paragraph.runs:
                run.bold = True
                run.font.size = Pt(16)
                run.font.color.rgb = RGBColor(31, 61, 105)
            break

    doc.core_properties.title = "Balaagh AI — User Interface Section"
    doc.core_properties.subject = "Completed dashboard and UI section for the Balaagh AI capstone prototype"
    doc.core_properties.comments = "Includes screenshots captured from the running Balaagh AI prototype."
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()